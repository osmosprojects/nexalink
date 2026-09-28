import os
import subprocess
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

excel_path = '/Users/saket/Downloads/NexaLink/nexalink_database_export.xlsx'
csv_dir = '/Users/saket/Downloads/NexaLink/database_csv_export'
os.makedirs(csv_dir, exist_ok=True)

# 1. Fetch tables list
cmd_tables = ['/Applications/XAMPP/xamppfiles/bin/mysql', '-u', 'root', '-B', '-e', 'USE nexalink_crm; SHOW TABLES;']
res = subprocess.run(cmd_tables, capture_output=True, text=True)
tables = [line.strip() for line in res.stdout.strip().split('\n')[1:] if line.strip()]

print(f"Exporting {len(tables)} tables to Excel & CSV...")

table_descriptions = {
    'users': 'Core application user accounts (email, display name, status, created timestamp)',
    'user_profiles': 'User professional profiles (headline, bio, company, job title, industry, skills, goals)',
    'user_personas': 'AI matchmaking personas (communication style, networking goals, target personas)',
    'contacts': 'CRM contact directory (relationship type, strength score, follow-up dates)',
    'tags': 'User-defined contact tags and categories',
    'contact_tags': 'Junction mapping between contacts and tags',
    'interactions': 'Timeline of touchpoints, meetings, notes, and calls',
    'meetings': 'Scheduled calendar meetings and event details',
    'tasks': 'Actionable tasks, due dates, and completion status',
    'notes': 'Saved user notes and meeting summaries',
    'goals': 'Long-term networking goals',
    'goal_progress': 'Progress tracking milestones for goals',
    'posts': 'Community feed posts',
    'recommendations': 'AI matchmaker recommendations',
    'skipped_profiles': 'Profiles skipped in AI discovery feed',
    'user_match_scores': 'Calculated AI compatibility match scores',
    'notifications': 'System & in-app user notifications',
    'audit_logs': 'Security and user action audit logs',
    'oauth_accounts': 'Linked third-party SSO accounts'
}

summary_rows = []

with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
    for tbl in tables:
        cmd_query = ['/Applications/XAMPP/xamppfiles/bin/mysql', '-u', 'root', '-B', '-e', f'USE nexalink_crm; SELECT * FROM `{tbl}`;']
        r = subprocess.run(cmd_query, capture_output=True, text=True)
        raw_lines = [line for line in r.stdout.split('\n') if line]
        
        if len(raw_lines) >= 1:
            headers = raw_lines[0].split('\t')
            if len(raw_lines) > 1:
                data = [line.split('\t') for line in raw_lines[1:]]
                df = pd.DataFrame(data, columns=headers)
            else:
                df = pd.DataFrame(columns=headers)
        else:
            df = pd.DataFrame()
            
        desc = table_descriptions.get(tbl, 'Database table')
        summary_rows.append({
            'Table Name': tbl,
            'Description': desc,
            'Total Rows': len(df),
            'Columns Count': len(df.columns)
        })
        
        # Save CSV
        csv_path = os.path.join(csv_dir, f'{tbl}.csv')
        df.to_csv(csv_path, index=False)
        print(f" - Exported {tbl}: {len(df)} rows, {len(df.columns)} columns")
        
        # Save Excel Sheet (sheet name max 31 chars)
        df.to_excel(writer, sheet_name=tbl[:31], index=False)
        
    summary_df = pd.DataFrame(summary_rows)
    summary_df.to_excel(writer, sheet_name='0_DB_Overview', index=False)

# 2. Styling Excel Workbook
wb = load_workbook(excel_path)
if '0_DB_Overview' in wb.sheetnames:
    overview_sheet = wb['0_DB_Overview']
    wb._sheets.remove(overview_sheet)
    wb._sheets.insert(0, overview_sheet)

header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid') # Navy Blue
header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
cell_font = Font(name='Calibri', size=10)
thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

for ws in wb.worksheets:
    ws.views.sheetView[0].showGridLines = True
    ws.row_dimensions[1].height = 26
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 1:
                cell.font = cell_font
                cell.border = thin_border
                cell.alignment = Alignment(vertical='center')
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 55)

wb.save(excel_path)
print("SUCCESS: Full database export completed successfully!")
